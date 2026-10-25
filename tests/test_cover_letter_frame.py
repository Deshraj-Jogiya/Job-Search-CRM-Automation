"""The cover letter's frame (greeting, date, contact line, sign-off), the Word and plain-text copies of both documents,
the addressee and referral form, and the rule that a cover letter's numbers must come from the candidate's own facts.
No real AI calls: the provider is patched everywhere."""

import io
import json
from datetime import date
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
import pdfplumber

from app.models import TailoredDocument
from app.services import cover_letter_service, document_render_service, tailoring_service

from tests.test_tailor_application_prose_checks import _BASE_PROFILE, _application

CONTACT = {"location": "Tempe, AZ", "zip": "85281", "phone": "555-0100", "email": "test@example.com",
           "linkedin": "https://linkedin.com/in/test"}
BODY = "First paragraph.\n\nSecond paragraph.\n\nThird paragraph."


class TestFrame:
    def test_greeting_uses_the_saved_name_or_hiring_manager(self):
        assert cover_letter_service.greeting("Jordan Lee") == "Dear Jordan Lee,"
        assert cover_letter_service.greeting(None) == "Dear Hiring Manager,"
        assert cover_letter_service.greeting("   ") == "Dear Hiring Manager,"

    def test_full_text_is_greeting_body_and_signoff_in_order(self):
        text = cover_letter_service.full_text(BODY, "Test Candidate")
        assert text.startswith("Dear Hiring Manager,\n\nFirst paragraph.")
        assert text.endswith("Third paragraph.\n\nSincerely,\nTest Candidate")

    def test_plain_text_has_contact_line_and_date(self):
        text = cover_letter_service.plain_text(BODY, "Test Candidate", CONTACT, "Jordan Lee", today=date(2026, 10, 25))
        lines = text.splitlines()
        assert lines[0] == "Test Candidate"
        assert lines[1] == "Tempe, AZ 85281 | 555-0100 | test@example.com | https://linkedin.com/in/test"
        assert "October 25, 2026" in lines and "Dear Jordan Lee," in lines

    def test_pdf_and_docx_carry_the_same_frame(self):
        pdf = document_render_service.render_cover_letter_pdf(BODY, "Test Candidate", CONTACT, None, today=date(2026, 10, 25))
        with pdfplumber.open(io.BytesIO(pdf)) as reader:
            pdf_text = "".join(p.extract_text() for p in reader.pages)
        for part in ("Test Candidate", "October 25, 2026", "Dear Hiring Manager,", "Second paragraph.", "Sincerely,"):
            assert part in pdf_text
        docx_text = [p.text for p in cover_letter_service.build_docx(BODY, "Test Candidate", CONTACT, "Jordan Lee").paragraphs]
        assert docx_text[0] == "Test Candidate" and "Dear Jordan Lee," in docx_text and docx_text[-2:] == ["Sincerely,", "Test Candidate"]

    def test_resume_plain_text_matches_the_word_document(self):
        resume = {**_BASE_PROFILE, "contact": CONTACT}
        text = cover_letter_service.resume_plain_text(resume)
        assert text.startswith("Test Candidate") and "PROFESSIONAL EXPERIENCE" in text and "- Built pipelines." in text


class TestNumbers:
    def test_a_number_not_in_the_facts_is_reported(self):
        facts = json.dumps({"bullets": ["Cut load time by 30% across 1,200 tables."]})
        assert tailoring_service.unsupported_numbers("Cut load time by 30% across 1200 tables.", facts) == []
        assert tailoring_service.unsupported_numbers("Saved 45% and served 10,000 users.", facts) == ["45%", "10,000"]


class TestPrompt:
    def _prompt(self, **kwargs):
        llm = MagicMock()
        llm.complete_text.return_value = "Body."
        with patch("app.services.tailoring_service.get_llm_provider", return_value=llm):
            tailoring_service.generate_cover_letter(_BASE_PROFILE, "Acme Corp", "Data Engineer", "Build pipelines.", **kwargs)
        return llm.complete_text.call_args.kwargs["prompt"]

    def test_names_the_role_uses_work_experience_and_asks_for_an_interview(self):
        prompt = self._prompt()
        assert "'Data Engineer'" in prompt and "Data Engineer at Real Co: Built pipelines." in prompt
        assert "interview" in prompt and "do not mention a referral" in prompt

    def test_a_saved_referral_is_named(self):
        prompt = self._prompt(referred_by="Sam Rivera")
        assert "referred by Sam Rivera" in prompt and "do not mention a referral" not in prompt


class TestTailoringRetriesAnInventedNumber:
    def _run(self, db, letters):
        application = _application(db, _BASE_PROFILE)
        generate = MagicMock(side_effect=letters)
        with (
            patch("app.services.tailoring_service.run_multi_pass_tailoring",
                  return_value=(_BASE_PROFILE["experience"], [], 90, [], [])),
            patch("app.services.tailoring_service._tailor_summary_skills", return_value={"summary": "Data engineer.", "skills": {}}),
            patch("app.services.tailoring_service.generate_cover_letter", generate),
            patch("app.services.tailoring_service.score_cover_letter", return_value=80),
        ):
            return tailoring_service.tailor_application(db, application.id), generate

    def test_a_clean_rewrite_is_kept(self, db):
        result, generate = self._run(db, ["I cut costs by 40%.", "I built pipelines."])
        assert generate.call_count == 2 and "40%" in generate.call_args.kwargs["correction"]
        assert result.attention_reason is None

    def test_a_number_still_invented_after_the_rewrite_needs_review(self, db):
        result, _ = self._run(db, ["I cut costs by 40%.", "I cut costs by 40%."])
        assert result.attention_reason and "cover letter number 40%" in result.attention_reason


@pytest.fixture()
def client(db, monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app as web_app

    monkeypatch.setenv("APP_MODE", "showcase")
    return TestClient(web_app)


def _tailored(db):
    application = _application(db, _BASE_PROFILE)
    db.add(TailoredDocument(application_id=application.id, document_type="resume",
                            content=json.dumps({**_BASE_PROFILE, "contact": CONTACT})))
    db.add(TailoredDocument(application_id=application.id, document_type="cover_letter", content=BODY))
    db.commit()
    return application


class TestRoutes:
    def test_word_and_plain_text_downloads(self, db, client):
        application = _tailored(db)
        letter = client.get(f"/jobs/{application.id}/tailored/cover_letter/download-docx")
        assert letter.status_code == 200 and letter.headers["content-disposition"].endswith('cover_letter.docx"')
        assert "Dear Hiring Manager," in [p.text for p in Document(io.BytesIO(letter.content)).paragraphs]
        txt = client.get(f"/jobs/{application.id}/tailored/cover_letter/download-txt")
        assert txt.headers["content-type"].startswith("text/plain") and "Sincerely,\nTest Candidate" in txt.text
        resume = client.get(f"/jobs/{application.id}/tailored/resume/download-txt")
        assert resume.status_code == 200 and "Built pipelines." in resume.text

    def test_saving_letter_details_changes_the_greeting(self, db, client):
        from app.csrf import CSRF_COOKIE_NAME

        application = _tailored(db)
        client.get("/how-scoring-works")
        response = client.post(f"/jobs/{application.id}/letter-details",
                               data={"recruiter_name": "Jordan Lee", "referred_by": "Sam Rivera",
                                     "csrf_token": client.cookies.get(CSRF_COOKIE_NAME)}, follow_redirects=False)
        assert response.status_code == 303 and "re-tailor" in response.headers["location"]
        db.refresh(application)
        assert (application.recruiter_name, application.referred_by) == ("Jordan Lee", "Sam Rivera")
        assert "Dear Jordan Lee," in client.get(f"/jobs/{application.id}/tailored/cover_letter/download-txt").text
